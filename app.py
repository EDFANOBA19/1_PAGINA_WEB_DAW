from flask import Flask, render_template, redirect, url_for, flash, request, send_file
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
import psycopg2
import psycopg2.extras
import io
from datetime import datetime

# Librerías para PDF
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_LEFT

# Librerías para Excel
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Importar formularios
from forms.producto_form import ProductoForm
from forms.cliente_form import ClienteForm
from forms.proveedor_form import ProveedorForm
from forms.facturacion_form import FacturaForm
from forms.abono_form import AbonoForm
from forms.login_form import LoginForm
from forms.registro_form import RegistroForm

# Importar conexión y modelo
from conexion.conexion import get_db_connection
from models import Usuario

app = Flask(__name__)

# ============================================================
# CONFIGURACIÓN
# ============================================================

app.config['SECRET_KEY'] = 'maquirenthal-secret-key-2026'

app.config['DB_HOST'] = 'localhost'
app.config['DB_NAME'] = 'maquirenthal_db'
app.config['DB_USER'] = 'postgres'
app.config['DB_PASSWORD'] = 'postgres123'
app.config['DB_PORT'] = '5432'

# ============================================================
# CONFIGURACIÓN DE LOGIN
# ============================================================

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'
login_manager.login_message = '⚠️ Debes iniciar sesión para acceder a esta página.'
login_manager.login_message_category = 'warning'

@login_manager.user_loader
def load_user(user_id):
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM usuarios WHERE id = %s', (user_id,))
    data = cursor.fetchone()
    cursor.close()
    conn.close()
    if data:
        return Usuario(data['id'], data['usuario'], data['email'], data['password'])
    return None

# ============================================================
# VARIABLES SIMPLES
# ============================================================

NOMBRE_EMPRESA = "maquirenthal"
ANIO_FUNDACION = 2023

# ============================================================
# FUNCIONES AUXILIARES
# ============================================================

def obtener_proveedores():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT id, nombre FROM proveedores ORDER BY nombre')
    proveedores = cursor.fetchall()
    cursor.close()
    conn.close()
    return proveedores

def obtener_clientes():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT id, nombre FROM clientes ORDER BY nombre')
    clientes = cursor.fetchall()
    cursor.close()
    conn.close()
    return clientes

def obtener_productos():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT id_producto, nombre, capacidad, precio FROM productos ORDER BY nombre')
    productos = cursor.fetchall()
    cursor.close()
    conn.close()
    return productos

def obtener_siguiente_numero_factura():
    try:
        conn = get_db_connection()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cursor.execute('SELECT COUNT(*) as total FROM facturas')
        total = cursor.fetchone()['total']
        siguiente = f'FAC-{str(total + 1).zfill(3)}'
        while True:
            cursor.execute('SELECT id FROM facturas WHERE numero = %s', (siguiente,))
            if not cursor.fetchone():
                break
            total += 1
            siguiente = f'FAC-{str(total + 1).zfill(3)}'
        cursor.close()
        conn.close()
        return siguiente
    except Exception as e:
        print(f"Error al generar número: {str(e)}")
        return 'FAC-001'

def obtener_estadisticas():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    
    cursor.execute('SELECT COUNT(*) as total FROM productos')
    total_productos = cursor.fetchone()['total']
    
    cursor.execute('SELECT COUNT(*) as total FROM clientes')
    total_clientes = cursor.fetchone()['total']
    
    cursor.execute('SELECT COUNT(*) as total FROM proveedores')
    total_proveedores = cursor.fetchone()['total']
    
    cursor.execute('SELECT COUNT(*) as total, COALESCE(SUM(total), 0) as monto FROM facturas')
    facturas_data = cursor.fetchone()
    total_facturas = facturas_data['total']
    monto_facturas = facturas_data['monto']
    
    cursor.execute('SELECT COUNT(*) as total FROM mensajes_contacto')
    total_mensajes = cursor.fetchone()['total']
    
    cursor.execute('SELECT COUNT(*) as total FROM mensajes_contacto WHERE leido = FALSE')
    mensajes_nuevos = cursor.fetchone()['total']
    
    cursor.execute('SELECT COUNT(*) as total FROM mensajes_contacto WHERE leido = TRUE')
    mensajes_leidos = cursor.fetchone()['total']
    
    cursor.execute('SELECT COUNT(*) as total FROM solicitudes_cotizacion')
    total_solicitudes = cursor.fetchone()['total']
    
    cursor.execute('SELECT COUNT(*) as total FROM solicitudes_cotizacion WHERE estado = %s', ('Pendiente',))
    solicitudes_pendientes = cursor.fetchone()['total']
    
    cursor.execute('SELECT COUNT(*) as total FROM solicitudes_cotizacion WHERE estado = %s', ('Atendida',))
    solicitudes_atendidas = cursor.fetchone()['total']
    
    cursor.close()
    conn.close()
    
    return {
        'total_productos': total_productos,
        'total_clientes': total_clientes,
        'total_proveedores': total_proveedores,
        'total_facturas': total_facturas,
        'monto_facturas': monto_facturas,
        'total_mensajes': total_mensajes,
        'mensajes_nuevos': mensajes_nuevos,
        'mensajes_leidos': mensajes_leidos,
        'total_solicitudes': total_solicitudes,
        'solicitudes_pendientes': solicitudes_pendientes,
        'solicitudes_atendidas': solicitudes_atendidas
    }

def generar_pdf(titulo, columnas, datos, nombre_archivo):
    """Genera un PDF con los datos proporcionados"""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    
    elementos = []
    estilos = getSampleStyleSheet()
    
    # Estilo del título
    estilo_titulo = ParagraphStyle(
        'Titulo',
        parent=estilos['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#0d47a1'),
        alignment=TA_CENTER,
        spaceAfter=10
    )
    
    # Estilo del subtítulo
    estilo_subtitulo = ParagraphStyle(
        'Subtitulo',
        parent=estilos['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#666666'),
        alignment=TA_CENTER,
        spaceAfter=20
    )
    
    # Encabezado
    elementos.append(Paragraph("MAQUIRENTHAL S.A.S.", estilo_titulo))
    elementos.append(Paragraph(f"RUC: 2293534872001 | {titulo}", estilo_subtitulo))
    elementos.append(Paragraph(f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}", estilo_subtitulo))
    elementos.append(Spacer(1, 20))
    
    # Tabla
    data = [columnas] + datos
    tabla = Table(data, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d47a1')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f0f0')]),
    ]))
    elementos.append(tabla)
    
    # Pie de página
    elementos.append(Spacer(1, 30))
    elementos.append(Paragraph("Gracias por confiar en Maquirenthal S.A.S.", estilo_subtitulo))
    
    doc.build(elementos)
    buffer.seek(0)
    return buffer

def generar_excel(titulo, columnas, datos, nombre_archivo):
    """Genera un archivo Excel con los datos proporcionados"""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:30]
    
    # Estilos
    font_titulo = Font(bold=True, size=14, color='FFFFFF')
    font_header = Font(bold=True, color='FFFFFF')
    fill_titulo = PatternFill(start_color='0D47A1', end_color='0D47A1', fill_type='solid')
    fill_header = PatternFill(start_color='1976D2', end_color='1976D2', fill_type='solid')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Título
    ws.merge_cells('A1:' + chr(64 + len(columnas)) + '1')
    ws['A1'] = f"MAQUIRENTHAL S.A.S. - {titulo}"
    ws['A1'].font = font_titulo
    ws['A1'].fill = fill_titulo
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    # Fecha
    ws.merge_cells('A2:' + chr(64 + len(columnas)) + '2')
    ws['A2'] = f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws['A2'].alignment = Alignment(horizontal='center')
    
    # Encabezados
    for col_num, columna in enumerate(columnas, 1):
        cell = ws.cell(row=4, column=col_num, value=columna)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Datos
    for row_num, fila in enumerate(datos, 5):
        for col_num, valor in enumerate(fila, 1):
            cell = ws.cell(row=row_num, column=col_num, value=valor)
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='center')
    
    # Ajustar ancho de columnas
    for col_num in range(1, len(columnas) + 1):
        ws.column_dimensions[chr(64 + col_num)].width = 20
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

# ============================================================
# RUTAS DE AUTENTICACIÓN
# ============================================================

@app.route('/registro', methods=['GET', 'POST'])
def registro():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    
    form = RegistroForm()
    if form.validate_on_submit():
        conn = get_db_connection()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        
        cursor.execute('SELECT * FROM usuarios WHERE usuario = %s', (form.usuario.data,))
        if cursor.fetchone():
            flash('❌ El nombre de usuario ya está registrado', 'danger')
            cursor.close()
            conn.close()
            return redirect(url_for('registro'))
        
        cursor.execute('SELECT * FROM usuarios WHERE email = %s', (form.email.data,))
        if cursor.fetchone():
            flash('❌ El correo electrónico ya está registrado', 'danger')
            cursor.close()
            conn.close()
            return redirect(url_for('registro'))
        
        password_hash = generate_password_hash(form.password.data)
        cursor.execute('INSERT INTO usuarios (usuario, email, password) VALUES (%s, %s, %s)',
                      (form.usuario.data, form.email.data, password_hash))
        conn.commit()
        cursor.close()
        conn.close()
        
        flash('✅ Usuario registrado correctamente.', 'success')
        return redirect(url_for('login'))
    
    return render_template('registro.html', form=form, empresa=NOMBRE_EMPRESA)

@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    
    form = LoginForm()
    if form.validate_on_submit():
        conn = get_db_connection()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cursor.execute('SELECT * FROM usuarios WHERE email = %s OR usuario = %s', 
                      (form.email.data, form.email.data))
        data = cursor.fetchone()
        cursor.close()
        conn.close()
        
        if data and check_password_hash(data['password'], form.password.data):
            user = Usuario(data['id'], data['usuario'], data['email'], data['password'])
            login_user(user)
            flash(f'✅ Bienvenido, {user.usuario}', 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('❌ Credenciales incorrectas.', 'danger')
    
    return render_template('login.html', form=form, empresa=NOMBRE_EMPRESA)

@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('✅ Sesión cerrada correctamente', 'success')
    return redirect(url_for('login'))

# ============================================================
# RUTAS PRINCIPALES
# ============================================================

@app.route('/')
def index():
    try:
        conn = get_db_connection()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cursor.execute('SELECT * FROM servicios WHERE activo = TRUE ORDER BY id')
        servicios = cursor.fetchall()
        cursor.close()
        conn.close()
    except Exception as e:
        print(f"❌ Error: {str(e)}")
        servicios = []
    
    return render_template('index.html', 
                          empresa=NOMBRE_EMPRESA.upper(),
                          anio=ANIO_FUNDACION,
                          servicios=servicios)

@app.route('/servicio/<int:id>')
def servicio_detalle(id):
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM servicios WHERE id = %s AND activo = TRUE', (id,))
    servicio = cursor.fetchone()
    cursor.close()
    conn.close()
    
    if not servicio:
        flash('❌ Servicio no encontrado', 'danger')
        return redirect(url_for('index') + '#servicios')
    
    return render_template('servicio_detalle.html', servicio=servicio, empresa=NOMBRE_EMPRESA.upper())

@app.route('/dashboard')
@login_required
def dashboard():
    estadisticas = obtener_estadisticas()
    return render_template('dashboard.html', empresa=NOMBRE_EMPRESA, stats=estadisticas)

# ============================================================
# RUTAS DE CONTACTO Y SOLICITUDES
# ============================================================

@app.route('/contacto', methods=['POST'])
def contacto():
    nombre = request.form.get('contactoNombre')
    email = request.form.get('contactoEmail')
    asunto = request.form.get('contactoAsunto')
    mensaje = request.form.get('contactoMensaje')
    
    if not all([nombre, email, asunto, mensaje]):
        flash('❌ Todos los campos son obligatorios', 'danger')
        return redirect(url_for('index') + '#contacto')
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('INSERT INTO mensajes_contacto (nombre, email, asunto, mensaje) VALUES (%s, %s, %s, %s)',
                      (nombre, email, asunto, mensaje))
        conn.commit()
        cursor.close()
        conn.close()
        flash('✅ ¡Mensaje enviado!', 'success')
    except Exception as e:
        flash(f'❌ Error: {str(e)}', 'danger')
    
    return redirect(url_for('index') + '#contacto')

@app.route('/solicitud', methods=['POST'])
def solicitud():
    nombre = request.form.get('solicitudNombre')
    email = request.form.get('solicitudEmail')
    categoria = request.form.get('solicitudCategoria')
    descripcion = request.form.get('solicitudDescripcion')
    
    if not all([nombre, email, categoria, descripcion]):
        flash('❌ Todos los campos son obligatorios', 'danger')
        return redirect(url_for('index') + '#solicitudes')
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('INSERT INTO solicitudes_cotizacion (nombre, email, categoria, descripcion) VALUES (%s, %s, %s, %s)',
                      (nombre, email, categoria, descripcion))
        conn.commit()
        cursor.close()
        conn.close()
        flash('✅ ¡Solicitud registrada!', 'success')
    except Exception as e:
        flash(f'❌ Error: {str(e)}', 'danger')
    
    return redirect(url_for('index') + '#solicitudes')

# ============================================================
# RUTAS DEL MÓDULO MENSAJES
# ============================================================

@app.route('/mensajes')
@login_required
def mensajes():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    
    cursor.execute('SELECT * FROM mensajes_contacto ORDER BY fecha DESC')
    mensajes_contacto = cursor.fetchall()
    
    cursor.execute('SELECT * FROM solicitudes_cotizacion ORDER BY fecha DESC')
    solicitudes = cursor.fetchall()
    
    for m in mensajes_contacto:
        cursor.execute('''SELECT * FROM respuestas_mensajes WHERE id_mensaje = %s AND tipo = 'contacto' ORDER BY fecha DESC''', (m['id'],))
        m['respuestas'] = cursor.fetchall()
    
    for s in solicitudes:
        cursor.execute('''SELECT * FROM respuestas_mensajes WHERE id_mensaje = %s AND tipo = 'solicitud' ORDER BY fecha DESC''', (s['id'],))
        s['respuestas'] = cursor.fetchall()
    
    cursor.close()
    conn.close()
    
    return render_template('mensajes.html', mensajes_contacto=mensajes_contacto, solicitudes=solicitudes, empresa=NOMBRE_EMPRESA)

@app.route('/mensaje/eliminar/<int:id>')
@login_required
def mensaje_eliminar(id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM respuestas_mensajes WHERE id_mensaje = %s AND tipo = %s', (id, 'contacto'))
    cursor.execute('DELETE FROM mensajes_contacto WHERE id = %s', (id,))
    conn.commit()
    cursor.close()
    conn.close()
    flash('✅ Mensaje eliminado', 'success')
    return redirect(url_for('mensajes') + '#contacto')

@app.route('/mensaje/leido/<int:id>')
@login_required
def mensaje_leido(id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('UPDATE mensajes_contacto SET leido = TRUE WHERE id = %s', (id,))
    conn.commit()
    cursor.close()
    conn.close()
    flash('✅ Mensaje marcado como leído', 'success')
    return redirect(url_for('mensajes') + '#contacto')

@app.route('/solicitud/eliminar/<int:id>')
@login_required
def solicitud_eliminar(id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM respuestas_mensajes WHERE id_mensaje = %s AND tipo = %s', (id, 'solicitud'))
    cursor.execute('DELETE FROM solicitudes_cotizacion WHERE id = %s', (id,))
    conn.commit()
    cursor.close()
    conn.close()
    flash('✅ Solicitud eliminada', 'success')
    return redirect(url_for('mensajes') + '#solicitudes')

@app.route('/solicitud/estado/<int:id>/<estado>')
@login_required
def solicitud_estado(id, estado):
    if estado not in ['Pendiente', 'Atendida', 'Cancelada']:
        flash('❌ Estado no válido', 'danger')
        return redirect(url_for('mensajes') + '#solicitudes')
    
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('UPDATE solicitudes_cotizacion SET estado = %s WHERE id = %s', (estado, id))
    conn.commit()
    cursor.close()
    conn.close()
    flash(f'✅ Estado: {estado}', 'success')
    return redirect(url_for('mensajes') + '#solicitudes')

@app.route('/responder/<tipo>/<int:id>', methods=['POST'])
@login_required
def responder(tipo, id):
    if tipo not in ['contacto', 'solicitud']:
        flash('❌ Tipo no válido', 'danger')
        return redirect(url_for('mensajes'))
    
    respuesta = request.form.get('respuesta')
    
    if not respuesta or len(respuesta.strip()) < 5:
        flash('❌ La respuesta debe tener al menos 5 caracteres', 'danger')
        return redirect(url_for('mensajes') + f'#{tipo}')
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('INSERT INTO respuestas_mensajes (id_mensaje, tipo, respuesta) VALUES (%s, %s, %s)', (id, tipo, respuesta.strip()))
        
        if tipo == 'contacto':
            cursor.execute('UPDATE mensajes_contacto SET leido = TRUE WHERE id = %s', (id,))
        elif tipo == 'solicitud':
            cursor.execute('UPDATE solicitudes_cotizacion SET estado = %s WHERE id = %s', ('Atendida', id))
        
        conn.commit()
        cursor.close()
        conn.close()
        flash('✅ Respuesta guardada.', 'success')
    except Exception as e:
        flash(f'❌ Error: {str(e)}', 'danger')
    
    return redirect(url_for('mensajes') + f'#{tipo}')

@app.route('/respuesta/eliminar/<int:id>')
@login_required
def respuesta_eliminar(id):
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT tipo FROM respuestas_mensajes WHERE id = %s', (id,))
    data = cursor.fetchone()
    tipo = data['tipo'] if data else 'contacto'
    
    cursor.execute('DELETE FROM respuestas_mensajes WHERE id = %s', (id,))
    conn.commit()
    cursor.close()
    conn.close()
    flash('✅ Respuesta eliminada', 'success')
    return redirect(url_for('mensajes') + f'#{tipo}')

# ============================================================
# CRUD - PRODUCTOS CON EXPORTACIÓN
# ============================================================

@app.route('/productos')
@login_required
def productos():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('''SELECT p.*, pr.nombre AS proveedor_nombre FROM productos p LEFT JOIN proveedores pr ON p.id_proveedor = pr.id ORDER BY p.id_producto DESC''')
    equipos = cursor.fetchall()
    cursor.close()
    conn.close()
    return render_template('productos.html', equipos=equipos, empresa=NOMBRE_EMPRESA)

@app.route('/productos/exportar/pdf')
@login_required
def productos_pdf():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('''SELECT p.*, pr.nombre AS proveedor_nombre FROM productos p LEFT JOIN proveedores pr ON p.id_proveedor = pr.id ORDER BY p.id_producto DESC''')
    equipos = cursor.fetchall()
    cursor.close()
    conn.close()
    
    columnas = ['ID', 'Nombre', 'Capacidad', 'Precio', 'Stock', 'Categoría', 'Proveedor']
    datos = []
    for eq in equipos:
        datos.append([
            str(eq['id_producto']),
            eq['nombre'],
            eq['capacidad'],
            f"${eq['precio']:.2f}",
            str(eq['stock']),
            eq['categoria'] or 'N/A',
            eq['proveedor_nombre'] or 'N/A'
        ])
    
    buffer = generar_pdf('LISTADO DE PRODUCTOS', columnas, datos, 'productos')
    return send_file(buffer, as_attachment=True, download_name=f'productos_{datetime.now().strftime("%Y%m%d")}.pdf', mimetype='application/pdf')

@app.route('/productos/exportar/excel')
@login_required
def productos_excel():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('''SELECT p.*, pr.nombre AS proveedor_nombre FROM productos p LEFT JOIN proveedores pr ON p.id_proveedor = pr.id ORDER BY p.id_producto DESC''')
    equipos = cursor.fetchall()
    cursor.close()
    conn.close()
    
    columnas = ['ID', 'Nombre', 'Capacidad', 'Precio', 'Stock', 'Categoría', 'Proveedor']
    datos = []
    for eq in equipos:
        datos.append([
            eq['id_producto'],
            eq['nombre'],
            eq['capacidad'],
            float(eq['precio']),
            eq['stock'],
            eq['categoria'] or 'N/A',
            eq['proveedor_nombre'] or 'N/A'
        ])
    
    buffer = generar_excel('LISTADO DE PRODUCTOS', columnas, datos, 'productos')
    return send_file(buffer, as_attachment=True, download_name=f'productos_{datetime.now().strftime("%Y%m%d")}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/productos/nuevo', methods=['GET', 'POST'])
@login_required
def producto_nuevo():
    form = ProductoForm()
    form.id_proveedor.choices = [(p['id'], p['nombre']) for p in obtener_proveedores()]
    
    if form.validate_on_submit():
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''INSERT INTO productos (nombre, capacidad, precio, stock, categoria, id_proveedor) VALUES (%s, %s, %s, %s, %s, %s)''', 
                      (form.nombre.data, form.capacidad.data, form.precio.data, form.stock.data, form.categoria.data, form.id_proveedor.data))
        conn.commit()
        cursor.close()
        conn.close()
        flash('✅ Producto creado', 'success')
        return redirect(url_for('productos'))
    return render_template('formulario_producto.html', form=form, titulo='Nuevo Producto', empresa=NOMBRE_EMPRESA)

@app.route('/productos/editar/<int:id>', methods=['GET', 'POST'])
@login_required
def producto_editar(id):
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM productos WHERE id_producto = %s', (id,))
    producto = cursor.fetchone()
    cursor.close()
    conn.close()
    
    if not producto:
        flash('❌ Producto no encontrado', 'danger')
        return redirect(url_for('productos'))
    
    form = ProductoForm(data=producto)
    form.id_proveedor.choices = [(p['id'], p['nombre']) for p in obtener_proveedores()]
    
    if form.validate_on_submit():
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''UPDATE productos SET nombre = %s, capacidad = %s, precio = %s, stock = %s, categoria = %s, id_proveedor = %s WHERE id_producto = %s''',
                      (form.nombre.data, form.capacidad.data, form.precio.data, form.stock.data, form.categoria.data, form.id_proveedor.data, id))
        conn.commit()
        cursor.close()
        conn.close()
        flash('✅ Producto actualizado', 'success')
        return redirect(url_for('productos'))
    
    return render_template('formulario_producto.html', form=form, titulo='Editar Producto', empresa=NOMBRE_EMPRESA)

@app.route('/productos/eliminar/<int:id>')
@login_required
def producto_eliminar(id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM productos WHERE id_producto = %s', (id,))
    conn.commit()
    cursor.close()
    conn.close()
    flash('✅ Producto eliminado', 'success')
    return redirect(url_for('productos'))

# ============================================================
# CRUD - CLIENTES CON EXPORTACIÓN
# ============================================================

@app.route('/clientes')
@login_required
def clientes():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM clientes ORDER BY id DESC')
    clientes = cursor.fetchall()
    cursor.close()
    conn.close()
    return render_template('clientes.html', clientes=clientes, empresa=NOMBRE_EMPRESA)

@app.route('/clientes/exportar/pdf')
@login_required
def clientes_pdf():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM clientes ORDER BY id DESC')
    clientes = cursor.fetchall()
    cursor.close()
    conn.close()
    
    columnas = ['ID', 'Nombre', 'Empresa', 'Teléfono', 'Email', 'Ciudad']
    datos = []
    for c in clientes:
        datos.append([
            str(c['id']),
            c['nombre'],
            c['empresa'],
            c['telefono'],
            c['email'],
            c['ciudad'] or 'N/A'
        ])
    
    buffer = generar_pdf('LISTADO DE CLIENTES', columnas, datos, 'clientes')
    return send_file(buffer, as_attachment=True, download_name=f'clientes_{datetime.now().strftime("%Y%m%d")}.pdf', mimetype='application/pdf')

@app.route('/clientes/exportar/excel')
@login_required
def clientes_excel():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM clientes ORDER BY id DESC')
    clientes = cursor.fetchall()
    cursor.close()
    conn.close()
    
    columnas = ['ID', 'Nombre', 'Empresa', 'Teléfono', 'Email', 'Ciudad']
    datos = []
    for c in clientes:
        datos.append([
            c['id'],
            c['nombre'],
            c['empresa'],
            c['telefono'],
            c['email'],
            c['ciudad'] or 'N/A'
        ])
    
    buffer = generar_excel('LISTADO DE CLIENTES', columnas, datos, 'clientes')
    return send_file(buffer, as_attachment=True, download_name=f'clientes_{datetime.now().strftime("%Y%m%d")}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/clientes/nuevo', methods=['GET', 'POST'])
@login_required
def cliente_nuevo():
    form = ClienteForm()
    if form.validate_on_submit():
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''INSERT INTO clientes (nombre, empresa, telefono, email, ciudad) VALUES (%s, %s, %s, %s, %s)''',
                      (form.nombre.data, form.empresa.data, form.telefono.data, form.email.data, form.ciudad.data))
        conn.commit()
        cursor.close()
        conn.close()
        flash('✅ Cliente creado', 'success')
        return redirect(url_for('clientes'))
    return render_template('formulario_cliente.html', form=form, titulo='Nuevo Cliente', empresa=NOMBRE_EMPRESA)

@app.route('/clientes/editar/<int:id>', methods=['GET', 'POST'])
@login_required
def cliente_editar(id):
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM clientes WHERE id = %s', (id,))
    cliente = cursor.fetchone()
    cursor.close()
    conn.close()
    
    if not cliente:
        flash('❌ Cliente no encontrado', 'danger')
        return redirect(url_for('clientes'))
    
    form = ClienteForm(data=cliente)
    if form.validate_on_submit():
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''UPDATE clientes SET nombre = %s, empresa = %s, telefono = %s, email = %s, ciudad = %s WHERE id = %s''',
                      (form.nombre.data, form.empresa.data, form.telefono.data, form.email.data, form.ciudad.data, id))
        conn.commit()
        cursor.close()
        conn.close()
        flash('✅ Cliente actualizado', 'success')
        return redirect(url_for('clientes'))
    
    return render_template('formulario_cliente.html', form=form, titulo='Editar Cliente', empresa=NOMBRE_EMPRESA)

@app.route('/clientes/eliminar/<int:id>')
@login_required
def cliente_eliminar(id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM clientes WHERE id = %s', (id,))
    conn.commit()
    cursor.close()
    conn.close()
    flash('✅ Cliente eliminado', 'success')
    return redirect(url_for('clientes'))

# ============================================================
# CRUD - PROVEEDORES CON EXPORTACIÓN
# ============================================================

@app.route('/proveedores')
@login_required
def proveedores():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM proveedores ORDER BY id DESC')
    proveedores = cursor.fetchall()
    cursor.close()
    conn.close()
    return render_template('proveedores.html', proveedores=proveedores, empresa=NOMBRE_EMPRESA)

@app.route('/proveedores/exportar/pdf')
@login_required
def proveedores_pdf():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM proveedores ORDER BY id DESC')
    proveedores = cursor.fetchall()
    cursor.close()
    conn.close()
    
    columnas = ['ID', 'Nombre', 'Especialidad', 'Contacto', 'Teléfono', 'País']
    datos = []
    for p in proveedores:
        datos.append([
            str(p['id']),
            p['nombre'],
            p['especialidad'],
            p['contacto'],
            p['telefono'],
            p['pais'] or 'N/A'
        ])
    
    buffer = generar_pdf('LISTADO DE PROVEEDORES', columnas, datos, 'proveedores')
    return send_file(buffer, as_attachment=True, download_name=f'proveedores_{datetime.now().strftime("%Y%m%d")}.pdf', mimetype='application/pdf')

@app.route('/proveedores/exportar/excel')
@login_required
def proveedores_excel():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM proveedores ORDER BY id DESC')
    proveedores = cursor.fetchall()
    cursor.close()
    conn.close()
    
    columnas = ['ID', 'Nombre', 'Especialidad', 'Contacto', 'Teléfono', 'País']
    datos = []
    for p in proveedores:
        datos.append([
            p['id'],
            p['nombre'],
            p['especialidad'],
            p['contacto'],
            p['telefono'],
            p['pais'] or 'N/A'
        ])
    
    buffer = generar_excel('LISTADO DE PROVEEDORES', columnas, datos, 'proveedores')
    return send_file(buffer, as_attachment=True, download_name=f'proveedores_{datetime.now().strftime("%Y%m%d")}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/proveedores/nuevo', methods=['GET', 'POST'])
@login_required
def proveedor_nuevo():
    form = ProveedorForm()
    if form.validate_on_submit():
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''INSERT INTO proveedores (nombre, especialidad, contacto, telefono, pais) VALUES (%s, %s, %s, %s, %s)''',
                      (form.nombre.data, form.especialidad.data, form.contacto.data, form.telefono.data, form.pais.data))
        conn.commit()
        cursor.close()
        conn.close()
        flash('✅ Proveedor creado', 'success')
        return redirect(url_for('proveedores'))
    return render_template('formulario_proveedor.html', form=form, titulo='Nuevo Proveedor', empresa=NOMBRE_EMPRESA)

@app.route('/proveedores/editar/<int:id>', methods=['GET', 'POST'])
@login_required
def proveedor_editar(id):
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM proveedores WHERE id = %s', (id,))
    proveedor = cursor.fetchone()
    cursor.close()
    conn.close()
    
    if not proveedor:
        flash('❌ Proveedor no encontrado', 'danger')
        return redirect(url_for('proveedores'))
    
    form = ProveedorForm(data=proveedor)
    if form.validate_on_submit():
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''UPDATE proveedores SET nombre = %s, especialidad = %s, contacto = %s, telefono = %s, pais = %s WHERE id = %s''',
                      (form.nombre.data, form.especialidad.data, form.contacto.data, form.telefono.data, form.pais.data, id))
        conn.commit()
        cursor.close()
        conn.close()
        flash('✅ Proveedor actualizado', 'success')
        return redirect(url_for('proveedores'))
    
    return render_template('formulario_proveedor.html', form=form, titulo='Editar Proveedor', empresa=NOMBRE_EMPRESA)

@app.route('/proveedores/eliminar/<int:id>')
@login_required
def proveedor_eliminar(id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM proveedores WHERE id = %s', (id,))
    conn.commit()
    cursor.close()
    conn.close()
    flash('✅ Proveedor eliminado', 'success')
    return redirect(url_for('proveedores'))

# ============================================================
# CRUD - FACTURAS CON ABONOS Y EXPORTACIÓN
# ============================================================

@app.route('/facturacion')
@login_required
def facturacion():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('''
        SELECT f.*, c.nombre AS cliente_nombre, p.nombre AS producto_nombre
        FROM facturas f
        LEFT JOIN clientes c ON f.id_cliente = c.id
        LEFT JOIN productos p ON f.id_producto = p.id_producto
        ORDER BY f.id DESC
    ''')
    facturas = cursor.fetchall()
    cursor.close()
    conn.close()
    return render_template('facturacion.html', facturas=facturas, empresa=NOMBRE_EMPRESA)

@app.route('/facturacion/exportar/pdf')
@login_required
def facturacion_pdf():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('''
        SELECT f.*, c.nombre AS cliente_nombre
        FROM facturas f
        LEFT JOIN clientes c ON f.id_cliente = c.id
        ORDER BY f.id DESC
    ''')
    facturas = cursor.fetchall()
    cursor.close()
    conn.close()
    
    columnas = ['N° Factura', 'Cliente', 'Período', 'Días', 'Total', 'Abonado', 'Saldo', 'Estado']
    datos = []
    for f in facturas:
        datos.append([
            f['numero'],
            f['cliente_nombre'],
            f"{f['fecha_inicio']} al {f['fecha_fin']}",
            str(f['dias']),
            f"${f['total']:.2f}",
            f"${f['total_abonado']:.2f}",
            f"${f['saldo_pendiente']:.2f}",
            f['estado']
        ])
    
    buffer = generar_pdf('LISTADO DE FACTURAS', columnas, datos, 'facturas')
    return send_file(buffer, as_attachment=True, download_name=f'facturas_{datetime.now().strftime("%Y%m%d")}.pdf', mimetype='application/pdf')

@app.route('/facturacion/exportar/excel')
@login_required
def facturacion_excel():
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('''
        SELECT f.*, c.nombre AS cliente_nombre
        FROM facturas f
        LEFT JOIN clientes c ON f.id_cliente = c.id
        ORDER BY f.id DESC
    ''')
    facturas = cursor.fetchall()
    cursor.close()
    conn.close()
    
    columnas = ['N° Factura', 'Cliente', 'Período', 'Días', 'Total', 'Abonado', 'Saldo', 'Estado']
    datos = []
    for f in facturas:
        datos.append([
            f['numero'],
            f['cliente_nombre'],
            f"{f['fecha_inicio']} al {f['fecha_fin']}",
            f['dias'],
            float(f['total']),
            float(f['total_abonado']),
            float(f['saldo_pendiente']),
            f['estado']
        ])
    
    buffer = generar_excel('LISTADO DE FACTURAS', columnas, datos, 'facturas')
    return send_file(buffer, as_attachment=True, download_name=f'facturas_{datetime.now().strftime("%Y%m%d")}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/facturacion/nuevo', methods=['GET', 'POST'])
@login_required
def factura_nueva():
    form = FacturaForm()
    form.id_cliente.choices = [(c['id'], c['nombre']) for c in obtener_clientes()]
    form.id_producto.choices = [(p['id_producto'], f"{p['nombre']} - {p['capacidad']} - ${p['precio']}") for p in obtener_productos()]
    
    if not form.numero.data:
        form.numero.data = obtener_siguiente_numero_factura()
    
    if form.validate_on_submit():
        conn = get_db_connection()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cursor.execute('SELECT id FROM facturas WHERE numero = %s', (form.numero.data,))
        factura_existente = cursor.fetchone()
        
        if factura_existente:
            sugerido = obtener_siguiente_numero_factura()
            flash(f'⚠️ ERROR: La factura {form.numero.data} YA EXISTE. Sugerencia: {sugerido}', 'danger')
            cursor.close()
            conn.close()
            return redirect(url_for('factura_nueva'))
        
        dias = (form.fecha_fin.data - form.fecha_inicio.data).days
        if dias <= 0:
            flash('❌ La fecha de fin debe ser mayor a la fecha de inicio', 'danger')
            cursor.close()
            conn.close()
            return redirect(url_for('factura_nueva'))
        
        cursor.execute('SELECT precio FROM productos WHERE id_producto = %s', (form.id_producto.data,))
        producto = cursor.fetchone()
        tarifa_diaria = float(producto['precio'])
        
        subtotal = dias * tarifa_diaria
        iva = subtotal * 0.15
        total = subtotal + iva
        
        forma_pago = form.forma_pago.data
        meses_diferidos = 0
        tipo_tarjeta = None
        
        if forma_pago == 'Crédito Directo':
            meses_diferidos = int(request.form.get('meses_credito_input', 0) or 0)
        elif forma_pago == 'Tarjeta de Crédito':
            meses_diferidos = int(request.form.get('meses_tarjeta_input', 0) or 0)
            tipo_tarjeta = request.form.get('tipo_tarjeta_input', '') or None
        
        cuota_mensual = 0
        if meses_diferidos > 0:
            cuota_mensual = total / meses_diferidos
        
        try:
            cursor.execute('''
                INSERT INTO facturas (numero, id_cliente, id_producto, fecha_inicio, fecha_fin, 
                                      dias, tarifa_diaria, subtotal, iva, total, total_abonado, 
                                      saldo_pendiente, forma_pago, estado, observaciones,
                                      meses_diferidos, cuota_mensual, tipo_tarjeta)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                RETURNING id
            ''', (form.numero.data, form.id_cliente.data, form.id_producto.data,
                  form.fecha_inicio.data, form.fecha_fin.data, dias, tarifa_diaria,
                  subtotal, iva, total, 0, total, forma_pago, 'Pendiente', 
                  form.observaciones.data, meses_diferidos, cuota_mensual, tipo_tarjeta))
            
            factura_id = cursor.fetchone()['id']
            conn.commit()
            cursor.close()
            conn.close()
            
            flash(f'✅ Factura {form.numero.data} creada correctamente', 'success')
            return redirect(url_for('factura_detalle', id=factura_id))
            
        except psycopg2.errors.UniqueViolation:
            conn.rollback()
            cursor.close()
            conn.close()
            flash(f'⚠️ ERROR: La factura {form.numero.data} YA EXISTE.', 'danger')
            return redirect(url_for('factura_nueva'))
        except Exception as e:
            conn.rollback()
            cursor.close()
            conn.close()
            flash(f'❌ Error: {str(e)}', 'danger')
            return redirect(url_for('factura_nueva'))
    
    return render_template('formulario_facturacion.html', form=form, titulo='Nueva Factura', empresa=NOMBRE_EMPRESA)

@app.route('/facturacion/ver/<int:id>')
@login_required
def factura_detalle(id):
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('''
        SELECT f.*, c.nombre AS cliente_nombre, c.empresa AS cliente_empresa,
               c.telefono AS cliente_telefono, c.email AS cliente_email, c.ciudad AS cliente_ciudad,
               p.nombre AS producto_nombre, p.capacidad AS producto_capacidad
        FROM facturas f
        LEFT JOIN clientes c ON f.id_cliente = c.id
        LEFT JOIN productos p ON f.id_producto = p.id_producto
        WHERE f.id = %s
    ''', (id,))
    factura = cursor.fetchone()
    
    if not factura:
        flash('❌ Factura no encontrada', 'danger')
        cursor.close()
        conn.close()
        return redirect(url_for('facturacion'))
    
    cursor.execute('SELECT * FROM abonos WHERE id_factura = %s ORDER BY fecha_abono DESC', (id,))
    abonos = cursor.fetchall()
    
    cursor.close()
    conn.close()
    
    return render_template('factura_detalle.html', factura=factura, abonos=abonos, empresa=NOMBRE_EMPRESA)

@app.route('/facturacion/abono/<int:id>', methods=['GET', 'POST'])
@login_required
def abono_nuevo(id):
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute('SELECT * FROM facturas WHERE id = %s', (id,))
    factura = cursor.fetchone()
    cursor.close()
    conn.close()
    
    if not factura:
        flash('❌ Factura no encontrada', 'danger')
        return redirect(url_for('facturacion'))
    
    form = AbonoForm()
    
    if form.validate_on_submit():
        monto = float(form.monto.data)
        
        if monto > float(factura['saldo_pendiente']):
            flash(f'❌ El monto no puede ser mayor al saldo pendiente (${factura["saldo_pendiente"]})', 'danger')
            return redirect(url_for('abono_nuevo', id=id))
        
        nuevo_total_abonado = float(factura['total_abonado']) + monto
        nuevo_saldo = float(factura['total']) - nuevo_total_abonado
        
        if nuevo_saldo <= 0:
            nuevo_estado = 'Pagada'
            nuevo_saldo = 0
        else:
            nuevo_estado = 'Abonada'
        
        conn = get_db_connection()
        cursor = conn.cursor()
        
        cursor.execute('''
            INSERT INTO abonos (id_factura, monto, forma_pago, referencia, observaciones)
            VALUES (%s, %s, %s, %s, %s)
        ''', (id, monto, form.forma_pago.data, form.referencia.data, form.observaciones.data))
        
        cursor.execute('''
            UPDATE facturas 
            SET total_abonado = %s, saldo_pendiente = %s, estado = %s
            WHERE id = %s
        ''', (nuevo_total_abonado, nuevo_saldo, nuevo_estado, id))
        
        conn.commit()
        cursor.close()
        conn.close()
        
        flash(f'✅ Abono de ${monto:.2f} registrado. Nuevo saldo: ${nuevo_saldo:.2f}', 'success')
        return redirect(url_for('factura_detalle', id=id))
    
    return render_template('formulario_abono.html', form=form, factura=factura, empresa=NOMBRE_EMPRESA)

@app.route('/facturacion/eliminar/<int:id>')
@login_required
def factura_eliminar(id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('DELETE FROM abonos WHERE id_factura = %s', (id,))
    cursor.execute('DELETE FROM facturas WHERE id = %s', (id,))
    conn.commit()
    cursor.close()
    conn.close()
    flash('✅ Factura eliminada', 'success')
    return redirect(url_for('facturacion'))

# ============================================================
# EJECUCIÓN
# ============================================================

if __name__ == '__main__':
    app.run(debug=True)